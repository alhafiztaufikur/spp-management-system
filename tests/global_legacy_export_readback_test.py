"""Reopen every exported PDF/XLSX and independently check key totals and missing values."""
from pathlib import Path
from decimal import Decimal, ROUND_HALF_UP
import json
import re
import subprocess
import sys
from pypdf import PdfReader
from openpyxl import load_workbook

qa=Path(sys.argv[1]).resolve();artifacts=qa/'artifacts'
baseline=json.loads((qa/'reconciliation.json').read_text())['source'];checks=[]
def check(ok,label):
    assert ok,label
    checks.append(label);print('PASS',label,flush=True)
def cents(value):return int((Decimal(str(value))*100).quantize(Decimal('1'),rounding=ROUND_HALF_UP))
def money(value):
    whole,fraction=divmod(value,100)
    return 'Rp '+f'{whole:,}'.replace(',','.')+(f',{fraction:02d}' if fraction else '')
templates=['status','penerimaan','spp-tahunan','per-item','tabungan-siswa','saldo-tabungan','riwayat-tagihan','setoran','kas-tabungan']
for unit in [1,2,3,0]:
    for template in templates:
        stem=f'{unit}-{template}';pdf=PdfReader(artifacts/(stem+'.pdf'))
        text=' '.join(' '.join(page.extract_text() or '' for page in pdf.pages).split())
        check(len(pdf.pages)>0 and all(abs(min(float(p.mediabox.width),float(p.mediabox.height))-595.28)<2 for p in pdf.pages),stem+': PDF pages are A4')
        check('histori Legacy tanpa tanggal tidak termasuk total periode' not in text and 'kewajiban tidak tercatat' not in text,stem+': PDF has no obsolete repeated explanations')
        book=load_workbook(artifacts/(stem+'.xlsx'),data_only=True)
        values=[c.value for sheet in book for row in sheet for c in row if c.value is not None]
        check(bool(values) and any('NIS' in str(v) for v in values) if template not in ['setoran','kas-tabungan'] else bool(values),stem+': workbook can be reopened with report content')
        texts=' '.join(str(v) for v in values if isinstance(v,str))
        if template in ['status','per-item','spp-tahunan']:
            note='Total kewajiban hanya mencakup nominal yang diketahui'
            check(note in text and note in texts,stem+': partial obligation totals are explained once in both formats')
        if template in ['penerimaan','setoran','tabungan-siswa','kas-tabungan']:
            note='Termasuk transaksi tanpa tanggal sesuai periode kewajiban'
            check(note in text and note in texts,stem+': PDF and Excel retain the applied undated filter')
        if template=='setoran':
            units=[str(unit)] if unit else ['1','2','3'];expected=sum(baseline['cash'][u]+baseline['undated'][u] for u in units)
            check(money(expected) in text,stem+': PDF settlement equals independent source cents')
            check(expected in [cents(v) for v in values if isinstance(v,(int,float)) and not isinstance(v,bool)],stem+': Excel settlement is numeric and equals independent source cents')
        if template in ['status','per-item']:
            sheet=book.worksheets[0];columns=[]
            for row in sheet.iter_rows(min_row=10,max_row=11):
                for cell in row:
                    if str(cell.value or '').strip().lower() in ['tagihan','total tagihan']:columns.append(cell.column)
            check(bool(columns) and all(sheet.cell(12,col).value in [None,'—','-'] for col in columns),stem+': unknown principal is a missing value, not numeric zero')
        book.close()
# Render first/last pages of representative outputs using an independent PDF engine.
poppler=Path(sys.argv[2])
for stem in ['1-per-item','1-spp-tahunan','1-setoran','0-kas-tabungan']:
    reader=PdfReader(artifacts/(stem+'.pdf'))
    for page in sorted(set([1,len(reader.pages)])):
        target=artifacts/f'{stem}-page-{page}'
        subprocess.run([str(poppler),'-f',str(page),'-l',str(page),'-singlefile','-scale-to','1600','-png',str(artifacts/(stem+'.pdf')),str(target)],check=True,stdout=subprocess.DEVNULL,stderr=subprocess.PIPE)
        check(target.with_suffix('.png').stat().st_size>10000,f'{stem} page {page}: independently rendered for visual inspection')
(qa/'export-readback.json').write_text(json.dumps({'checks':len(checks),'passed':checks},indent=2))
