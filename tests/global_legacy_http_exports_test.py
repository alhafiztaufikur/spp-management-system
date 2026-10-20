"""Authenticated UI and binary exports on the explicitly owned migrated clone."""
from pathlib import Path
import json
import re
import sys
import time
import urllib.error
import urllib.parse
import urllib.request
from lxml import html

qa=Path(sys.argv[1]).resolve();f=json.loads((qa/'fixture.json').read_text())
assert f['base']=='http://127.0.0.1:58471' and f['port']==58470
checks=[];timings=[]
def check(ok,label):
    assert ok,label
    checks.append(label);print('PASS',label,flush=True)
def get(unit,path,params=None,role='sessions'):
    session=f[role][unit]
    url=f['base']+'/'+path+('?' + urllib.parse.urlencode(params,doseq=True) if params else '')
    req=urllib.request.Request(url,headers={'Cookie':'PHPSESSID='+session})
    started=time.perf_counter()
    try: response=urllib.request.urlopen(req,timeout=180)
    except urllib.error.HTTPError as e: response=e
    with response: body=response.read();status=response.status;headers=dict(response.headers)
    timings.append(dict(unit=unit,template=(params or {}).get('template'),format=(params or {}).get('format','html'),seconds=round(time.perf_counter()-started,3),bytes=len(body)))
    assert not any(x in body[:20000] for x in [b'Fatal error:',b'Warning:',b'Notice:',b'Uncaught ']),body[:1000]
    return status,body,headers
check(json.loads(get(1,'tests/browser_clone_identity.php')[1])['database']==f['database'],'HTTP endpoint uses the owned report clone')
if sys.argv[2:]==['--full-receipt']:
    from pypdf import PdfReader
    from openpyxl import load_workbook
    source=json.loads((qa/'reconciliation.json').read_text())['source']
    expected=sum(source['cash'].values())+sum(source['undated'].values())
    params=dict(template='penerimaan',unit='all',kategori=['*'],tanggal_awal='2026-10-01',tanggal_akhir='2026-10-21',siswa_status='all',include_undated='1',download='1')
    for format,extension in [('pdf','pdf'),('excel','xlsx')]:
        status,binary,_=get(0,'laporan/export_global.php',params|dict(format=format))
        check(status==200 and binary.startswith(b'%PDF-' if format=='pdf' else b'PK'),'Full all-unit '+format+' export is valid')
        (qa/'artifacts'/('0-penerimaan-full.'+extension)).write_bytes(binary)
    pdf=PdfReader(qa/'artifacts'/'0-penerimaan-full.pdf')
    page_text=[' '.join((page.extract_text() or '').split()) for page in pdf.pages];text=' '.join(page_text)
    check('Rp '+f'{expected//100:,}'.replace(',','.') in text,'Full PDF total equals independently summed source cents')
    book=load_workbook(qa/'artifacts'/'0-penerimaan-full.xlsx',data_only=True)
    check(any(round(cell.value*100)==expected for sheet in book for row in sheet for cell in row if isinstance(cell.value,(int,float))),'Full Excel total equals independently summed source cents')
    sheet=book.worksheets[0];header=next(row for row in sheet if any(cell.value=='NIS' for cell in row));nis_col=next(cell.column for cell in header if cell.value=='NIS');unit_col=next(cell.column for cell in header if cell.value=='Unit')
    records=[(int(sheet.cell(i,1).value),str(sheet.cell(i,unit_col).value),str(sheet.cell(i,nis_col).value)) for i in range(header[0].row+1,sheet.max_row+1) if isinstance(sheet.cell(i,1).value,(int,float))]
    check(bool(records) and all(re.search(r'\b'+str(number)+r'\s*'+re.escape(unit)+r'\s*'+re.escape(nis),text) for number,unit,nis in records),'Every Excel row identity and number is present in the complete PDF')
    check(all('Halaman '+str(i+1) in value for i,value in enumerate(page_text)),'Full PDF has continuous page numbers across render fragments')
    check(text.count('Termasuk transaksi tanpa tanggal sesuai periode kewajiban')==1,'Full PDF explains the undated filter once')
    book.close();(qa/'full-receipt.json').write_text(json.dumps(dict(checks=len(checks),passed=checks,timings=timings,pages=len(pdf.pages)),indent=2));sys.exit(0)

templates=['status','penerimaan','spp-tahunan','per-item','tabungan-siswa','saldo-tabungan','riwayat-tagihan','setoran','kas-tabungan']
for unit in [1,2,3,0]:
    nis=f['samples'][str(unit or 1)]['nis']
    for template in templates:
        params=dict(template=template,unit=str(unit) if unit else 'all',kategori='semua' if template=='penerimaan' else 'spp',tanggal_awal='2026-10-01',tanggal_akhir='2026-10-21',tahun_ajaran='2026/2027',bulan_awal='10',bulan_akhir='10',tahun_awal='2026',tahun_akhir='2026',siswa_status='all')
        if template not in ['setoran','kas-tabungan']:params['q']=nis
        if template in ['penerimaan','setoran','tabungan-siswa','kas-tabungan']:params['include_undated']='1'
        status,body,_=get(unit,'laporan/template.php',params)
        check(status==200 and b'Gagal memuat laporan' not in body,f'Unit {unit} {template} UI renders')
        check(b'histori Legacy tanpa tanggal tidak termasuk total periode' not in body and b'Histori Legacy \xc2\xb7 kewajiban tidak tercatat' not in body,f'Unit {unit} {template} does not repeat verbose Legacy labels')
        tree=html.fromstring(body)
        if 'include_undated' in params:
            box=tree.xpath('//input[@name="include_undated"]')
            check(len(box)==1 and box[0].get('checked') is not None,f'Unit {unit} {template} scoped undated checkbox is applied')
            export_links=tree.xpath('//a[contains(@class,"report-export-btn")]/@href')
            check(len(export_links)==3 and all('include_undated=1' in url for url in export_links),f'Unit {unit} {template} all export links retain the checkbox')
        for format,extension in [('pdf','pdf'),('excel','xlsx')]:
            status,binary,headers=get(unit,'laporan/export_global.php',params|dict(format=format,download='1'))
            check(status==200 and binary.startswith(b'%PDF-' if format=='pdf' else b'PK'),f'Unit {unit} {template} {format} is a valid binary document')
            (qa/'artifacts'/f'{unit}-{template}.{extension}').write_bytes(binary)
        status,body,_=get(unit,'laporan/template.php',params,'admin_sessions')
        check(status==200 and b'Gagal memuat laporan' not in body,f'Administrator reads unit {unit} {template}')
    # Unit switches preserve only safe, applicable query state; pagination always reuses it.
    params=dict(template='penerimaan',unit=str(unit) if unit else 'all',include_undated='1',tanggal_awal='2026-10-01',tanggal_akhir='2026-10-21',per_page='25',page='2')
    _,body,_=get(unit,'laporan/template.php',params);tree=html.fromstring(body)
    links=tree.xpath('//a[contains(@class,"global-unit-choice")]/@href')
    check(len(links)==4 and all('include_undated=1' in link for link in links),f'Unit {unit} scope links preserve the checkbox')
    page_links=tree.xpath('//a[contains(@href,"page=")]/@href')
    check(bool(page_links) and all('include_undated=1' in link for link in page_links),f'Unit {unit} pagination retains scoped undated transactions')
    # Test HTML print semantics as a PDF response, not a receipt or a partial table.
    status,binary,headers=get(unit,'laporan/export_global.php',params|dict(format='print',q=nis))
    check(status==200 and binary.startswith(b'%PDF-') and 'inline' in headers.get('Content-Disposition',''),f'Unit {unit} print opens the complete filtered PDF inline')
(qa/'http-exports.json').write_text(json.dumps({'checks':len(checks),'passed':checks,'timings':timings},indent=2))
